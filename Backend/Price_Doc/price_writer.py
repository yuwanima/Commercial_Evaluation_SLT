from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation



FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def _hdr(cell, text):
    cell.value = text
    cell.font = _font(bold=True, color="FFFFFF")
    cell.fill = HDR_FILL
    cell.border = BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="center")

AUTO = PatternFill("solid", fgColor="E8F0FE")        # light blue = formula
LOW = PatternFill("solid", fgColor="F8CBAD")         # orange = low-confidence OCR, please verify
YEL = PatternFill("solid", fgColor="FFF2CC")
BLUE = "0000FF"
LAST = 2000                                            # formulas look this far down so rows can be added

# (header, width, field key or None for formulas)
COLS = [
    ("File", 26, "file"), ("Supplier", 24, "supplier"), ("Part", 11, "part"), ("S/No", 7, "sno"),
    ("Description", 46, "desc"), ("Qty", 8, "qty"), ("Unit", 8, "unit"), ("Country of Origin", 13, "origin"),
    ("Currency", 9, "currency"), ("Unit Price (excl. taxes)", 15, "unit_excl"), ("VAT %", 7, "vat"), ("NBT %", 7, "nbt"),
    ("Other taxes %", 8, "other_tax"), ("Unit Price (incl. taxes) - as quoted", 16, "unit_incl"),
    ("Unit Price (incl. taxes) - calculated", 16, None), ("FOB/FCA Unit Price", 14, "fob_unit"), ("FOB/FCA Total", 14, "fob_total"),
    ("Insurance", 11, "insurance"), ("Freight", 11, "freight"), ("CIF/CIP Total", 14, "cif_total"),
    ("Agent's Commission", 12, "agent_comm"), ("Amount (as printed)", 15, "amount"), ("LINE TOTAL (calculated)", 16, None),
    ("Total based on", 26, None), ("Check", 34, None), ("Remarks", 36, "remarks"), ("Row type", 10, "rowtype"),
    ("Page", 6, "page"), ("Read by", 11, "method")]
CI = {key: get_column_letter(i + 1) for i, (_, _, key) in enumerate(COLS) if key}
NAMEOF = {h: get_column_letter(i + 1) for i, (h, _, _) in enumerate(COLS)}
LTOT, BASIS, CHECK, QINCL_CALC = NAMEOF["LINE TOTAL (calculated)"], NAMEOF["Total based on"], NAMEOF["Check"], NAMEOF["Unit Price (incl. taxes) - calculated"]
MONEY = {"unit_excl", "unit_incl", "fob_unit", "fob_total", "insurance", "freight", "cif_total", "agent_comm", "amount"}
NUMIN = MONEY | {"qty", "vat", "nbt", "other_tax"}
HDR_ROW, FIRST = 4, 5


def _cell(ws, r, c, v, bold=False, fmt=None, wrap=True, color=None, fill=None, center=False):
    cell = ws.cell(r, c, v)
    cell.font = _font(bold=bold, color=color) if color else _font(bold=bold)
    cell.border = BORDER
    cell.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal="center" if center else None)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def _formulas(r):
    c = CI
    f = {k: f"{v}{r}" for k, v in c.items()}
    N, J, K, L, M, F = f["unit_incl"], f["unit_excl"], f["vat"], f["nbt"], f["other_tax"], f["qty"]
    O = f"{QINCL_CALC}{r}"
    P, Q, R, S, T, U, V = f["fob_unit"], f["fob_total"], f["insurance"], f["freight"], f["cif_total"], f["agent_comm"], f["amount"]
    W, X = f"{LTOT}{r}", f"{BASIS}{r}"
    part, rtype, file_ = f["part"], f["rowtype"], f["file"]
    calc_incl = (f'=IF(AND(ISNUMBER({J}),OR(ISNUMBER({K}),ISNUMBER({L}),ISNUMBER({M}))),'
                 f'IF(Summary!$B$3="Sum of rates",{J}*(1+(N({K})+N({L})+N({M}))/100),'
                 f'{J}*(1+N({L})/100)*(1+N({K})/100)*(1+N({M})/100)),"")')
    total = (f'=IF({rtype}<>"Item","",IF({part}="B Foreign",IF(ISNUMBER({T}),{T},IF(ISNUMBER({V}),{V},"")),'
             f'IF(AND(ISNUMBER({F}),ISNUMBER({N})),{F}*{N},IF(AND(ISNUMBER({F}),ISNUMBER({O})),{F}*{O},'
             f'IF(AND(ISNUMBER({F}),ISNUMBER({J})),{F}*{J},IF(ISNUMBER({V}),{V},""))))))')
    basis = (f'=IF({W}="","",IF({part}="B Foreign",IF(ISNUMBER({T}),"CIF/CIP total (as quoted)","Amount as printed"),'
             f'IF(AND(ISNUMBER({F}),ISNUMBER({N})),"Qty x unit price incl. taxes (quoted)",'
             f'IF(AND(ISNUMBER({F}),ISNUMBER({O})),"Qty x unit price incl. taxes (calculated)",'
             f'IF(AND(ISNUMBER({F}),ISNUMBER({J})),"Qty x unit price (tax basis not stated)","Amount as printed")))))')
    rng = lambda col: f"${col}${FIRST}:${col}${LAST}"
    sum_items = (f'SUMIFS({rng(LTOT)},{rng(c["file"])},{file_},{rng(c["part"])},{part},{rng(c["rowtype"])},"Item")')
    check = (f'=IF({rtype}="Total",IF(ISNUMBER({V}),IF(ABS({V}-{sum_items})>0.01*ABS({V})+0.5,'
             f'"CHECK: printed total differs from sum of lines","OK"),""),'
             f'IF({rtype}<>"Item","",'
             f'IF(AND(ISNUMBER({N}),ISNUMBER({O}),ABS(N({N})-N({O}))>0.01*ABS(N({O}))),"CHECK: incl. price differs from excl. price + taxes",'
             f'IF(OR(AND({J}<>"",NOT(ISNUMBER({J}))),AND({N}<>"",NOT(ISNUMBER({N}))),AND({P}<>"",NOT(ISNUMBER({P}))),AND({T}<>"",NOT(ISNUMBER({T})))),"CHECK: price is not a number (e.g. Included / FOC)",'
             f'IF(AND(OR(ISNUMBER({J}),ISNUMBER({N}),ISNUMBER({P}),ISNUMBER({T})),NOT(ISNUMBER({F}))),"CHECK: quantity missing",'
             f'IF(AND(ISNUMBER({P}),ISNUMBER({Q}),ISNUMBER({F}),ABS(N({Q})-N({F})*N({P}))>0.01*ABS(N({Q}))+0.5),"CHECK: FOB total differs from qty x FOB unit",'
             f'IF(AND(ISNUMBER({T}),ISNUMBER({Q}),ABS(N({T})-N({Q})-N({R})-N({S}))>0.01*ABS(N({T}))+0.5),"CHECK: CIF differs from FOB + insurance + freight",'
             f'IF(AND(LEFT({X},3)="Qty",ISNUMBER({V}),ABS(N({V})-N({W}))>0.01*ABS(N({V}))+0.5),"CHECK: amount differs from qty x unit price",'
             f'IF(ISNUMBER({W}),"OK","")))))))))')
    return calc_incl, total, basis, check


def build_workbook(reports, out_path):
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws = wb.create_sheet("Standardized Lines")
    wr = wb.create_sheet("Raw Tables")
    wl = wb.create_sheet("Extraction Log")

    # ------------------------------------------------------------------ Standardized Lines
    ws["A1"] = "Standardized Lines - all price lines of all uploaded files (edit freely; formulas recalculate)"
    ws["A1"].font = _font(bold=True, size=13)
    ws["A2"] = ("Blue = typed/extracted number | Light-blue = formula | Orange = low-confidence OCR, please verify against the original | "
                f"'Check' compares quoted and calculated values (1 percent tolerance). Add rows below the last line - formulas look down to row {LAST}.")
    ws["A2"].font = _font(italic=True, size=9, color="7F6000")
    ws.sheet_view.showGridLines = False
    for i, (h, w, key) in enumerate(COLS, 1):
        _hdr(ws.cell(HDR_ROW, i), h)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[HDR_ROW].height = 44
    dv_part = DataValidation(type="list", formula1='"A Local,B Foreign,Other"', allow_blank=True, showErrorMessage=False)
    dv_type = DataValidation(type="list", formula1='"Item,Sub-total,Total,Note"', allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv_part)
    ws.add_data_validation(dv_type)
    r = FIRST
    per_file = {}
    for rep in reports:
        per_file[rep["source"]] = {"first": r}
        for ln in rep["lines"]:
            vals = dict(ln, file=rep["source"], supplier="")
            for i, (h, w, key) in enumerate(COLS, 1):
                col = get_column_letter(i)
                if key == "supplier":
                    f = f'=IFERROR(INDEX(Summary!$B$5:$B$500,MATCH({CI["file"]}{r},Summary!$A$5:$A$500,0)),{CI["file"]}{r})'
                    _cell(ws, r, i, f, fill=AUTO)
                elif key:
                    v = vals.get(key)
                    if key == "method":
                        v = {"docx": "Word", "pdf": "PDF text", "ocr": "OCR"}.get(v, v)
                    if v == "" or v is None:
                        v = None
                    cell = _cell(ws, r, i, v, fmt="#,##0.00" if key in MONEY else None,
                                 color=BLUE if key in NUMIN and isinstance(v, (int, float)) else None,
                                 center=key in ("sno", "page", "rowtype", "currency", "part"))
                    if key in ln.get("lowconf", ()):
                        cell.fill = LOW
                    if key == "part":
                        dv_part.add(cell)
                    if key == "rowtype":
                        dv_type.add(cell)
            calc, total, basis, check = _formulas(r)
            _cell(ws, r, [h for h, _, _ in COLS].index("Unit Price (incl. taxes) - calculated") + 1, calc, fmt="#,##0.00", fill=AUTO)
            _cell(ws, r, [h for h, _, _ in COLS].index("LINE TOTAL (calculated)") + 1, total, bold=True, fmt="#,##0.00", fill=AUTO)
            _cell(ws, r, [h for h, _, _ in COLS].index("Total based on") + 1, basis, fill=AUTO)
            _cell(ws, r, [h for h, _, _ in COLS].index("Check") + 1, check, fill=AUTO)
            r += 1
        per_file[rep["source"]]["last"] = r - 1
    last_data = max(r - 1, FIRST)
    ws.freeze_panes = ws.cell(FIRST, 6)
    ws.auto_filter.ref = f"A{HDR_ROW}:{get_column_letter(len(COLS))}{last_data}"
    red = Font(name=FONT, bold=True, color="C00000")
    grn = Font(name=FONT, color="2E7D32")
    ws.conditional_formatting.add(f"{CHECK}{FIRST}:{CHECK}{LAST}", FormulaRule(formula=[f'LEFT({CHECK}{FIRST},5)="CHECK"'], font=red))
    ws.conditional_formatting.add(f"{CHECK}{FIRST}:{CHECK}{LAST}", FormulaRule(formula=[f'{CHECK}{FIRST}="OK"'], font=grn))
    tot_font = Font(name=FONT, bold=True)
    ws.conditional_formatting.add(f"A{FIRST}:{get_column_letter(len(COLS))}{LAST}",
                                  FormulaRule(formula=[f'OR(${CI["rowtype"]}{FIRST}="Total",${CI["rowtype"]}{FIRST}="Sub-total")'], font=tot_font))

    # ------------------------------------------------------------------ Summary
    ws_sum["A1"] = "Price Schedule Summary"
    ws_sum["A1"].font = _font(bold=True, size=14)
    ws_sum["A2"] = "Supplier names can be edited here (column B) - they flow into 'Standardized Lines'. Totals are formulas over the Standardized Lines sheet."
    ws_sum["A2"].font = _font(italic=True, size=9, color="7F6000")
    ws_sum.sheet_view.showGridLines = False
    ws_sum["A3"] = "Tax method for 'calculated' unit price:"
    ws_sum["A3"].font = _font(bold=True)
    ws_sum["B3"] = "NBT first, then VAT"
    ws_sum["B3"].font = _font(color=BLUE, bold=True)
    ws_sum["B3"].fill = YEL
    ws_sum["B3"].border = BORDER
    dv_tax = DataValidation(type="list", formula1='"NBT first then VAT,Sum of rates"', allow_blank=False, showErrorMessage=True)
    ws_sum.add_data_validation(dv_tax)
    dv_tax.add(ws_sum["B3"])
    ws_sum["B3"].value = "NBT first then VAT"
    ws_sum["C3"] = ("Used only to cross-check quoted prices and when a quote shows no tax-inclusive price. "
                    "'NBT first then VAT' = price x (1+NBT) x (1+VAT) x (1+other); 'Sum of rates' = price x (1+(VAT+NBT+other)). "
                    "Choose the method that matches the tender conditions.")
    ws_sum["C3"].font = _font(italic=True, size=9, color="7F6000")
    heads = ["File", "Supplier (editable)", "Tender No", "Read by", "Pages", "Item lines", "Total - Local (A)", "Total - Foreign (B, CIF/CIP)",
             "Total - Other", "Total rows printed in document", "Lines flagged CHECK", "Low-confidence OCR cells", "Notes"]
    widths = [28, 28, 22, 18, 7, 10, 18, 20, 16, 30, 12, 14, 70]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        _hdr(ws_sum.cell(4, i), h)
        ws_sum.column_dimensions[get_column_letter(i)].width = w
    ws_sum.row_dimensions[4].height = 44
    SL = "'Standardized Lines'!"
    R = lambda col: f"{SL}${col}${FIRST}:${col}${LAST}"
    for k, rep in enumerate(reports):
        row = 5 + k
        fl = rep["fields"]
        items = [l for l in rep["lines"] if l["rowtype"] == "Item"]
        printed = []
        for part in ("A Local", "B Foreign", "Other"):
            t = [l for l in rep["lines"] if l["rowtype"] == "Total" and l["part"] == part and isinstance(l.get("amount"), float)]
            if t:
                printed.append(f"{part[0]}: {t[-1]['amount']:,.2f} {t[-1]['currency']}".strip())
        nlow = sum(len(l["lowconf"]) for l in rep["lines"])
        sup = fl["supplier"]
        vals = [rep["source"], sup, fl["tender_no"] or None, fl["method"], fl["pages"], None, None, None, None,
                "; ".join(printed) or None, None, nlow, "\n".join(rep["notes"]) or None]
        for i, v in enumerate(vals, 1):
            _cell(ws_sum, row, i, v)
        ws_sum.cell(row, 2).fill = YEL
        ws_sum.cell(row, 2).font = _font(color=BLUE)
        a = f"$A{row}"
        _cell(ws_sum, row, 6, f'=COUNTIFS({R(CI["file"])},{a},{R(CI["rowtype"])},"Item")', fill=AUTO, center=True)
        for col, part in ((7, "A Local"), (8, "B Foreign"), (9, "Other")):
            _cell(ws_sum, row, col, f'=SUMIFS({R(LTOT)},{R(CI["file"])},{a},{R(CI["part"])},"{part}",{R(CI["rowtype"])},"Item")',
                  fmt="#,##0.00", fill=AUTO, bold=True)
        _cell(ws_sum, row, 11, f'=COUNTIFS({R(CI["file"])},{a},{R(CHECK)},"CHECK*")', fill=AUTO, center=True)
        ws_sum.row_dimensions[row].height = max(30, 15 * (1 + len(rep["notes"])))
    if not reports:
        ws_sum["A5"] = "No files processed."
    ws_sum.conditional_formatting.add(f"K5:K{4 + max(len(reports), 1)}", FormulaRule(formula=["K5>0"], font=red))
    ws_sum.freeze_panes = "C5"

    # ------------------------------------------------------------------ Raw Tables
    wr["A1"] = "Raw Tables - every table exactly as read (verbatim). Orange = low-confidence OCR."
    wr["A1"].font = _font(bold=True, size=13)
    wr.sheet_view.showGridLines = False
    rr = 3
    maxc = 1
    for rep in reports:
        for t in rep["raw_tables"]:
            c = wr.cell(rr, 1, f"{rep['source']}  |  table {t['idx']}  |  page {t['page']}  |  read by {t['method']}  |  recognised as: {t['part']}  |  {t['n_lines']} line(s)")
            c.font = _font(bold=True, color="FFFFFF")
            c.fill = HDR_FILL
            ncol = max((len(x) for x in t["rows"]), default=1)
            for j in range(2, ncol + 1):
                wr.cell(rr, j).fill = HDR_FILL
            rr += 1
            for ri, row in enumerate(t["rows"]):
                for ci, v in enumerate(row):
                    cell = _cell(wr, rr, ci + 1, v or None, bold=(ri == 0))
                    cf = t.get("conf")
                    if cf and ri < len(cf) and ci < len(cf[ri]) and cf[ri][ci] is not None and cf[ri][ci] < 85:
                        cell.fill = LOW
                rr += 1
                maxc = max(maxc, len(row))
            rr += 1
    for j in range(1, maxc + 1):
        wr.column_dimensions[get_column_letter(j)].width = 26
    wr.column_dimensions["B"].width = 44

    # ------------------------------------------------------------------ Log
    wl["A1"] = "Extraction Log"
    wl["A1"].font = _font(bold=True, size=13)
    for j, h in enumerate(("File", "Read by", "Tables", "Item lines", "Notes"), 1):
        _hdr(wl.cell(3, j), h)
    for i, rep in enumerate(reports):
        row = [rep["source"], rep["fields"]["method"], len(rep["raw_tables"]),
               len([l for l in rep["lines"] if l["rowtype"] == "Item"]), "\n".join(rep["notes"]) or "-"]
        if rep["fields"]["doc_remarks"]:
            row[4] += "\nDocument remarks found:\n" + rep["fields"]["doc_remarks"]
        for j, v in enumerate(row, 1):
            _cell(wl, 4 + i, j, v)
    for col, w in zip("ABCDE", (30, 20, 8, 10, 110)):
        wl.column_dimensions[col].width = w

    for sh in wb.worksheets:
        sh.sheet_properties.tabColor = "B45F06" if sh.title in ("Summary", "Standardized Lines") else "7F7F7F"
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return out_path
