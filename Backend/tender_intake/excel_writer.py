
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
LBL_FILL = PatternFill("solid", fgColor="EDEDED")
BLANK_FILL = PatternFill("solid", fgColor="FFF2CC")   # yellow = nothing extracted, please fill in
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (field key, row label, hint shown as cell comment, dropdown list or None)
T1 = [("tender_no", "Tender No", "", None), ("user_department", "User Department", "", None),
      ("forwarded_by", "Forwarded by", "", None), ("_budget", "Budget Confirmed", "CAPEX / OPEX amounts", None),
      ("_jtb", "Previous JTB Reference (if any)", "JTB Ref No and Date", None),
      ("published", "Tender Published Date", "", None), ("closed", "Tender Closed Date", "", None),
      ("opened", "Technical Bid Opened Date", "", None), ("submitted", "SSET Report Submission Date", "", None),
      ("bond_expiry", "Bid Bond Expiry Date(s)", "", None), ("purpose", "Purpose of Tender (sec 2a)", "", None),
      ("description", "Short Description of Goods/Services (Sec 2b)",
       "Include details on installation, commissioning, training and the time period covered", None)]
T2 = [("need", "How the need arose", "", None),
      ("approval", "Relevant Approval", "Governance Board / DPAC approval etc.", None),
      ("project", "Project No. & Name", "", None),
      ("account", "Account Code & Cost Center", "if under OPEX", None),
      ("budget_res", "Budget Reserved/Total Available", "", None),
      ("method", "Method of Tendering", "Open Tender / Registered Suppliers / Shortlisted / Sole Source", None),
      ("process", "Procurement Process up to Bid Opening",
       "Events and dates: floating, bidders' conference, extensions, closing, opening", None),
      ("changes", "Changes to Tender Conditions", "Document any changes and attach relevant approvals", None)]
T3 = [("agent", "Local Agent"), ("loa", "Letter of Authority (Yes/No)"), ("brand", "Product Brand and Model"),
      ("validity", "Bid Validity (as mentioned in bid)"), ("bond", "Expiry Date of Bid Bond")]
CN = ["Compliant", "Non-compliant"]
T4 = [("c_general", "Compliance to General Conditions", "Compliant / Non-compliant", CN),
      ("c_technical", "Compliance to Technical Conditions", "Compliant / Non-compliant", CN),
      ("c_special", "Compliance to Special Conditions", "Compliant / Non-compliant", CN),
      ("c_contract", "Compliance to Contract Conditions", "Compliant / Non-compliant", CN),
      ("major", "Major Non-compliances", "List major non-compliances beyond acceptable levels that lead to disqualification", None),
      ("impact", "Impact of Non-compliances", "Detail the impact of the major non-compliances", None),
      ("minor", "Minor / Acceptable Non-compliances", "List minor non-compliances where the bidder still qualifies", None),
      ("negotiate", "Commercial Terms to be Negotiated",
       "Non-compliances to financial terms where the bid cannot be disqualified", None)]
T5 = [("recommended", "Recommended for Commercial Evaluation?", "Yes / No", ["Yes", "No"]),
      ("bom", "Evaluated BOM Provided?", "Yes / No - soft copy in Excel format required", ["Yes", "No"]),
      ("breakdown", "Breakdown of Services",
       "Are training, pre-shipment inspection, spare parts, services (installation, testing, commissioning) indicated with quantities?", None),
      ("amc", "AMC Status", "Included / 'AMC is not required for this project/Procurement'", None),
      ("warranty", "Warranty Period Included", "State the warranty period", None),
      ("missing", "Missing Prices for Items?", "If yes, detail the items requiring pricing in a separate sealed envelope", None),
      ("prev_prices", "Previous Purchase Prices", "Provide time period and quantity. If none, provide market reference prices", None),
      ("other", "Other Commercial Items", "Any other items accounted for commercial evaluation", None)]


def _font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def _title(ws, text, ncols):
    ws["A1"] = text
    ws["A1"].font = _font(bold=True, size=13)
    ws["A2"] = "Yellow cells = nothing was found in the report - please fill in. All cells are editable."
    ws["A2"].font = _font(italic=True, color="7F6000", size=9)
    ws.sheet_view.showGridLines = False


def _hdr(cell, text):
    cell.value = text
    cell.font = _font(bold=True, color="FFFFFF")
    cell.fill = HDR_FILL
    cell.border = BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="center")


def _is_empty(v):
    if v in (None, ""):
        return True
    return isinstance(v, str) and v.strip() in ("CAPEX: -\nOPEX: -", "JTB Ref No:\nDate:")


def _put(ws, r, c, value, hint="", dv=None, bold=False):
    cell = ws.cell(r, c)
    cell.value = value or None          # keep the CAPEX/OPEX and JTB skeleton text like the template
    cell.font = _font(bold=bold)
    cell.border = BORDER
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    if _is_empty(value):
        cell.fill = BLANK_FILL
    if hint and c > 1:
        cell.comment = Comment(hint, "SSET tool", width=260, height=70)
    if dv is not None:
        dv.add(cell)
    return cell


def _label(ws, r, text):
    c = ws.cell(r, 1, text)
    c.font = _font(bold=True)
    c.fill = LBL_FILL
    c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical="top")


def _src(rep):
    return f"{rep['fields'].get('tender_no') or rep['source']}\n[{rep['source']}]"


def _kv_sheet(wb, name, title, rows, reports):
    ws = wb.create_sheet(name)
    _title(ws, title, 1 + len(reports))
    _hdr(ws.cell(4, 1), "FIELD")
    for j, rep in enumerate(reports):
        _hdr(ws.cell(4, 2 + j), "DETAILS\n" + _src(rep))
    ws.row_dimensions[4].height = 42
    for i, (key, label, hint, options) in enumerate(rows):
        r = 5 + i
        _label(ws, r, label)
        dv = None
        if options:
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv)
        for j, rep in enumerate(reports):
            _put(ws, r, 2 + j, rep["fields"].get(key, ""), hint, dv)
    ws.column_dimensions["A"].width = 38
    for j in range(len(reports)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 55
    ws.freeze_panes = "B5"
    return ws


def build_workbook(reports, out_path):
    """reports: list of dicts {source, fields, bidders, notes}"""
    wb = Workbook()
    wb.remove(wb.active)

    _kv_sheet(wb, "01 Tender & Timeline", "01 Tender & Timeline Header", T1, reports)
    _kv_sheet(wb, "02 Background Summary", "02 Background Summary", T2, reports)

    # ---- 03 Bidder details: one column per bidder ----
    ws = wb.create_sheet("03 Bidder Details")
    _title(ws, "3. Bidder Details Table", 2)
    cols = [(rep, b) for rep in reports for b in (rep["bidders"] or [{}])]
    _hdr(ws.cell(4, 1), "FIELD")
    _label(ws, 5, "Bidder / Company Name")
    for j, (rep, b) in enumerate(cols):
        _hdr(ws.cell(4, 2 + j), "BIDDER / COMPANY NAME:")
        _put(ws, 5, 2 + j, b.get("name", ""), "(company name)", bold=True)
        for i, (key, label) in enumerate(T3):
            if j == 0:
                _label(ws, 6 + i, label)
            _put(ws, 6 + i, 2 + j, b.get(key, ""))
        if j == 0:
            _label(ws, 6 + len(T3), "Source report")
        c = _put(ws, 6 + len(T3), 2 + j, _src(rep).replace("\n", " "))
        c.fill = NOTE_FILL
        c.font = _font(italic=True, color="595959", size=9)
        ws.column_dimensions[get_column_letter(2 + j)].width = 36
    ws.column_dimensions["A"].width = 38
    ws.freeze_panes = "B5"

    _kv_sheet(wb, "04 Compliance & Observations", "4. Compliance & Observations Table", T4, reports)
    _kv_sheet(wb, "05 Commercial Evaluation", "5. Recommendations for Commercial Evaluation", T5, reports)

    # ---- extraction log ----
    ws = wb.create_sheet("Extraction Log")
    ws["A1"] = "Extraction Log (for the reviewer - not part of the 5 tables)"
    ws["A1"].font = _font(bold=True, size=13)
    ws.sheet_view.showGridLines = False
    _hdr(ws.cell(3, 1), "Source file")
    _hdr(ws.cell(3, 2), "Fields found")
    _hdr(ws.cell(3, 3), "Fields missing")
    _hdr(ws.cell(3, 4), "Notes")
    wanted = [(k, l) for k, l, *_ in T1 + T2 + T4 + T5 if not k.startswith("_")]
    for i, rep in enumerate(reports):
        miss = [l for k, l in wanted if k not in rep["fields"]]
        notes = list(rep["notes"])
        if not rep["bidders"]:
            notes.append("No bidder table found / bidder table empty.")
        row = [rep["source"], f"{len(wanted) - len(miss)} of {len(wanted)}", "; ".join(miss) or "-", "\n".join(notes) or "-"]
        for j, v in enumerate(row):
            c = ws.cell(4 + i, 1 + j, v)
            c.font = _font()
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", (34, 14, 70, 60)):
        ws.column_dimensions[col].width = w

    for w in wb.worksheets:
        w.sheet_properties.tabColor = "1F3864" if w.title[0].isdigit() else "7F7F7F"
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out_path)
    return out_path
